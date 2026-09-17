import hashlib
from pathlib import Path

from openpyxl import load_workbook
from ontome_importer.cli import main
from ontome_importer.profiles import load_audit_profiles, load_generation_profiles, verify_capability_xsd


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "fixtures/phase2/rdf/baseline.rdf"


def test_init_creates_a_valid_auditable_workspace(tmp_path, capsys):
    workspace = tmp_path / "my-import"
    assert main([
        "init", "--source", str(SOURCE), "--workspace", str(workspace),
        "--scope-uri-prefix", "https://example.org/",
    ]) == 0
    assert "Next command:" in capsys.readouterr().out
    assert {path.name for path in workspace.iterdir()} == {"README.md", "source", "config", "build"}
    assert "Prochaine étape : audit" in (workspace / "README.md").read_text()
    copied = workspace / "source/ontology.rdf"
    assert copied.read_bytes() == SOURCE.read_bytes()
    manifest = (workspace / "config/audit.yaml").read_text()
    assert hashlib.sha256(copied.read_bytes()).hexdigest() in manifest
    assert "file: ../source/ontology.rdf" in manifest
    profiles = load_audit_profiles(workspace / "config/audit.yaml")
    verify_capability_xsd(profiles.capability)
    generation = load_generation_profiles(workspace / "config/generation.yaml")
    verify_capability_xsd(generation.capability)
    assert main(["audit", "--manifest", str(workspace / "config/audit.yaml"), "--output-dir", str(workspace / "build/audit")]) == 0
    assert {path.name for path in (workspace / "build/audit").iterdir()} == {"inventory.json", "audit.json", "audit.md"}


def test_init_renders_multiple_scope_prefixes_and_refuses_existing_workspace(tmp_path, capsys):
    workspace = tmp_path / "my-import"
    arguments = [
        "init", "--source", str(SOURCE), "--format", "rdfxml", "--workspace", str(workspace),
        "--scope-uri-prefix", "https://example.org/one/", "--scope-uri-prefix", "https://example.org/two/",
    ]
    assert main(arguments) == 0
    mapping = (workspace / "config/profiles/mapping-audit.yaml").read_text()
    assert "https://example.org/one/" in mapping
    assert "https://example.org/two/" in mapping
    assert main(arguments) == 2
    assert "Workspace already exists" in capsys.readouterr().err


def test_init_requires_an_explicit_format_when_extension_is_unknown(tmp_path, capsys):
    source = tmp_path / "ontology.data"
    source.write_bytes(SOURCE.read_bytes())
    workspace = tmp_path / "my-import"
    assert main(["init", "--source", str(source), "--workspace", str(workspace), "--scope-uri-prefix", "https://example.org/"]) == 2
    assert "Cannot infer RDF format" in capsys.readouterr().err
    assert not workspace.exists()


def test_audit_creates_a_mapping_workbook_for_a_workspace(tmp_path):
    workspace = tmp_path / "my-import"
    assert main([
        "init", "--source", str(SOURCE), "--workspace", str(workspace), "--scope-uri-prefix", "https://example.org/",
    ]) == 0
    workbook = workspace / "decisions/mapping.xlsx"
    assert main([
        "audit", "--manifest", str(workspace / "config/audit.yaml"), "--generation-manifest", str(workspace / "config/generation.yaml"),
        "--output-dir", str(workspace / "build/audit"), "--workbook", str(workbook),
    ]) == 0
    document = load_workbook(workbook, read_only=True)
    assert {"SUMMARY", "CLASSES", "PROPERTIES", "EXTERNAL_REFERENCES", "METADATA", "BLOCKERS", "VALIDATION"} <= set(document.sheetnames)
