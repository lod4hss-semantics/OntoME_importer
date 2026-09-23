from pathlib import Path
import shutil

from openpyxl import load_workbook
import pytest
import yaml

from ontome_importer.cli import main
from ontome_importer.loader import load_inventory
from ontome_importer.mapping_assistant import AssistantError, _catalog_identifiers, _external_uris
from ontome_importer.profiles import load_generation_profiles


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "fixtures/phase4/import-manifest.yaml"


def test_assistant_round_trips_existing_generation_profiles(tmp_path):
    workbook = tmp_path / "mapping.xlsx"
    check = tmp_path / "check.json"
    mapping = tmp_path / "mapping.yaml"
    registry = tmp_path / "registry.yaml"
    report = tmp_path / "compile.json"
    assert main(["assist", "export", "--manifest", str(MANIFEST), "--output", str(workbook)]) == 0
    assert main(["assist", "check", "--manifest", str(MANIFEST), "--workbook", str(workbook), "--output", str(check)]) == 0
    assert main([
        "assist", "compile", "--manifest", str(MANIFEST), "--workbook", str(workbook),
        "--mapping-output", str(mapping), "--registry-output", str(registry), "--report-output", str(report),
    ]) == 0
    original_mapping = yaml.safe_load((ROOT / "fixtures/phase4/profiles/mapping.yaml").read_text())
    assert yaml.safe_load(mapping.read_text()) == original_mapping
    assert yaml.safe_load(registry.read_text()) == yaml.safe_load((ROOT / "fixtures/phase4/profiles/namespace-registry.yaml").read_text())
    assert tuple(cell.value for cell in load_workbook(workbook)["EXTERNAL_REFERENCE_RULES"][1]) == ("id", "uri_prefix", "reference_namespace", "identifier_source", "identifier_pattern")
    assert tuple(cell.value for cell in load_workbook(workbook)["DECISIONS"][1]) == ("id", "resource_uri", "construct", "action", "status", "rationale", "approved_by", "approved_at", "decision_reference", "mapping_rule", "exception_id")


def test_assistant_rejects_a_stale_workbook(tmp_path, capsys):
    workbook = tmp_path / "mapping.xlsx"
    assert main(["assist", "export", "--manifest", str(MANIFEST), "--output", str(workbook)]) == 0
    document = load_workbook(workbook)
    metadata = document["_metadata"]
    metadata["B2"] = "not-the-source-checksum"
    document.save(workbook)
    assert main(["assist", "check", "--manifest", str(MANIFEST), "--workbook", str(workbook), "--output", str(tmp_path / "check.json")]) == 3
    assert "source checksum" not in capsys.readouterr().err


def test_assistant_round_trips_generic_external_reference_rules(tmp_path):
    workspace = tmp_path / "workspace"
    shutil.copytree(ROOT / "fixtures/phase4", workspace)
    profile = workspace / "profiles/mapping.yaml"
    mapping = yaml.safe_load(profile.read_text())
    mapping["external_references"] = []
    mapping["external_reference_rules"] = [{"id": "external-suffix", "uri_prefix": "https://example.org/external/", "reference_namespace": 100, "identifier_extraction": {"source": "uri_suffix"}}]
    profile.write_text(yaml.safe_dump(mapping, sort_keys=False), encoding="utf-8")
    workbook = tmp_path / "mapping.xlsx"
    output = tmp_path / "mapping.yaml"
    assert main(["assist", "export", "--manifest", str(workspace / "import-manifest.yaml"), "--output", str(workbook)]) == 0
    document = load_workbook(workbook)
    assert document["EXTERNAL_EXCEPTIONS"].max_row == 1
    assert {row[4] for row in document["EXTERNAL_USAGE"].iter_rows(min_row=2, values_only=True)} == {"resolved"}
    assert {row[5] for row in document["EXTERNAL_USAGE"].iter_rows(min_row=2, values_only=True)} == {"rule"}
    assert main(["assist", "compile", "--manifest", str(workspace / "import-manifest.yaml"), "--workbook", str(workbook), "--mapping-output", str(output), "--registry-output", str(tmp_path / "registry.yaml"), "--report-output", str(tmp_path / "report.json")]) == 0
    assert yaml.safe_load(output.read_text()) == mapping


def test_check_is_read_only_and_refresh_writes_derived_validation(tmp_path):
    workbook = tmp_path / "mapping.xlsx"
    assert main(["assist", "export", "--manifest", str(MANIFEST), "--output", str(workbook)]) == 0
    original = workbook.read_bytes()
    assert main(["assist", "check", "--manifest", str(MANIFEST), "--workbook", str(workbook), "--output", str(tmp_path / "check.json")]) == 0
    assert workbook.read_bytes() == original
    assert main(["assist", "refresh", "--manifest", str(MANIFEST), "--workbook", str(workbook), "--output", str(tmp_path / "refresh.json")]) == 0
    validation = load_workbook(workbook)["VALIDATION"]
    assert tuple(cell.value for cell in validation[1]) == ("phase", "severity", "sheet", "row", "resource_uri", "mapping_rule", "external_uri", "triple_ids", "code", "message")


def test_compile_rejects_colliding_outputs_without_modifying_existing_file(tmp_path):
    workbook = tmp_path / "mapping.xlsx"
    output = tmp_path / "output.yaml"
    output.write_text("original\n", encoding="utf-8")
    assert main(["assist", "export", "--manifest", str(MANIFEST), "--output", str(workbook)]) == 0
    assert main([
        "assist", "compile", "--manifest", str(MANIFEST), "--workbook", str(workbook),
        "--mapping-output", str(output), "--registry-output", str(output), "--report-output", str(tmp_path / "report.json"),
    ]) == 2
    assert output.read_text(encoding="utf-8") == "original\n"


def test_compile_preserves_unreferenced_namespace_rows(tmp_path):
    workbook = tmp_path / "mapping.xlsx"
    assert main(["assist", "export", "--manifest", str(MANIFEST), "--output", str(workbook)]) == 0
    document = load_workbook(workbook)
    document["NAMESPACE_REGISTRY"].append(("https://example.org/reserved/", 999, "forbidden", "test", ""))
    document.save(workbook)
    registry = tmp_path / "registry.yaml"
    assert main([
        "assist", "compile", "--manifest", str(MANIFEST), "--workbook", str(workbook),
        "--mapping-output", str(tmp_path / "mapping.yaml"), "--registry-output", str(registry), "--report-output", str(tmp_path / "report.json"),
    ]) == 0
    assert any(item["ontome_namespace_id"] == 999 for item in yaml.safe_load(registry.read_text())["namespaces"])


def test_compile_rejects_invalid_rule_json_before_publishing(tmp_path):
    workbook = tmp_path / "mapping.xlsx"
    mapping = tmp_path / "mapping.yaml"
    mapping.write_text("original\n", encoding="utf-8")
    assert main(["assist", "export", "--manifest", str(MANIFEST), "--output", str(workbook)]) == 0
    document = load_workbook(workbook)
    document["RULES"]["N2"] = "{not-json}"
    document.save(workbook)
    assert main([
        "assist", "compile", "--manifest", str(MANIFEST), "--workbook", str(workbook),
        "--mapping-output", str(mapping), "--registry-output", str(tmp_path / "registry.yaml"), "--report-output", str(tmp_path / "report.json"),
    ]) == 2
    assert mapping.read_text(encoding="utf-8") == "original\n"


def test_compile_does_not_publish_any_artifact_when_a_destination_is_unusable(tmp_path):
    workbook = tmp_path / "mapping.xlsx"
    mapping = tmp_path / "mapping.yaml"
    registry = tmp_path / "registry.yaml"
    mapping.write_text("old mapping\n", encoding="utf-8")
    registry.write_text("old registry\n", encoding="utf-8")
    blocker = tmp_path / "not-a-directory"
    blocker.write_text("file\n", encoding="utf-8")
    assert main(["assist", "export", "--manifest", str(MANIFEST), "--output", str(workbook)]) == 0
    assert main([
        "assist", "compile", "--manifest", str(MANIFEST), "--workbook", str(workbook),
        "--mapping-output", str(mapping), "--registry-output", str(registry), "--report-output", str(blocker / "report.json"),
    ]) == 2
    assert mapping.read_text(encoding="utf-8") == "old mapping\n"
    assert registry.read_text(encoding="utf-8") == "old registry\n"


def test_catalog_identifiers_require_one_literal_per_uri(tmp_path):
    catalog = tmp_path / "catalog.nt"
    predicate = "https://example.org/identifier"
    catalog.write_text(
        '<https://example.org/external/Term> <https://example.org/identifier> "T1" .\n',
        encoding="utf-8",
    )
    inventory = load_inventory(catalog, "ntriples")
    assert _catalog_identifiers(inventory, predicate) == {"https://example.org/external/Term": "T1"}
    catalog.write_text(
        '<https://example.org/external/Term> <https://example.org/identifier> "T1" .\n'
        '<https://example.org/external/Term> <https://example.org/identifier> "T2" .\n',
        encoding="utf-8",
    )
    with pytest.raises(AssistantError, match="ambiguous"):
        _catalog_identifiers(load_inventory(catalog, "ntriples"), predicate)


def test_external_reference_list_respects_rdf_type_scope_selectors():
    profiles = load_generation_profiles(MANIFEST)
    inventory = load_inventory(ROOT / "fixtures/phase4/rdf/valid.ttl", "turtle")
    mapping = dict(profiles.mapping)
    mapping["scope"] = {"resource_selectors": [{"rdf_type": "http://www.w3.org/2002/07/owl#Class"}]}
    assert "https://example.org/source/Parent" not in _external_uris(inventory, mapping)
