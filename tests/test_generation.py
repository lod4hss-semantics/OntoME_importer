import json
from pathlib import Path

from jsonschema import Draft202012Validator
from lxml import etree
from openpyxl import load_workbook

from ontome_importer.cli import main
from ontome_importer.loader import load_inventory
from ontome_importer.profiles import load_generation_profiles, verify_capability_xsd
from ontome_importer.resolution import resolve_generation
from ontome_importer.xml_writer import write_xml


ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "fixtures/phase4"


def test_generation_resolves_writes_valid_xml_and_complete_trace(tmp_path):
    profiles = load_generation_profiles(FIXTURES / "import-manifest.yaml")
    inventory = load_inventory(FIXTURES / "rdf/valid.ttl", "turtle")
    result = resolve_generation(inventory, profiles.manifest, profiles.capability, profiles.mapping, profiles.namespace_registry)
    assert result.generation is not None
    xml, trace = write_xml(result.generation, profiles.capability, verify_capability_xsd(profiles.capability), inventory.source_sha256)
    document = etree.fromstring(xml)
    assert document.xpath("/namespace/classes/class[1]/equivalentClass/@referenceNamespace") == ["100"]
    assert document.xpath("/namespace/properties/property[identifierInNamespace='property']/hasDomain/text()") == ["Class"]
    assert document.xpath("/namespace/properties/property[identifierInNamespace='dataProperty']/hasRange/text()") == ["D1"]
    assert document.xpath("/namespace/properties/property[identifierInNamespace='dataProperty']/hasRange/@referenceNamespace") == ["100"]
    assert document.xpath("/namespace/properties/property/identifierInNamespace/text()") == ["dataProperty", "property", "rdfProperty"]
    assert document.xpath("/namespace/referenceNamespace/text()") == ["100"]
    assert document.xpath("/namespace/classes/class[1]/textProperties/contextNote/text()") == ["A mapped comment."]
    assert all(entry["source_triples"] for entry in trace["entries"] if entry["origin"] == "rdf")
    trace_schema = json.loads((ROOT / "schemas/reports/generation-trace-1.2.schema.json").read_text())
    audit_schema = json.loads((ROOT / "schemas/reports/generation-audit-1.1.schema.json").read_text())
    assert list(Draft202012Validator(trace_schema).iter_errors(trace)) == []
    assert list(Draft202012Validator(audit_schema).iter_errors(result.audit)) == []


def test_generation_is_byte_deterministic(tmp_path):
    profiles = load_generation_profiles(FIXTURES / "import-manifest.yaml")
    inventory = load_inventory(FIXTURES / "rdf/valid.ttl", "turtle")
    first = resolve_generation(inventory, profiles.manifest, profiles.capability, profiles.mapping, profiles.namespace_registry)
    second = resolve_generation(inventory, profiles.manifest, profiles.capability, profiles.mapping, profiles.namespace_registry)
    first_xml, first_trace = write_xml(first.generation, profiles.capability, verify_capability_xsd(profiles.capability), inventory.source_sha256)
    second_xml, second_trace = write_xml(second.generation, profiles.capability, verify_capability_xsd(profiles.capability), inventory.source_sha256)
    assert first_xml == second_xml
    assert first_trace == second_trace


def test_missing_property_range_blocks_resolution(tmp_path):
    profiles = load_generation_profiles(FIXTURES / "import-manifest.yaml")
    source = tmp_path / "missing-range.ttl"
    source.write_text((FIXTURES / "rdf/valid.ttl").read_text().replace("  rdfs:range ex:Parent ;\n", ""))
    inventory = load_inventory(source, "turtle")
    result = resolve_generation(inventory, profiles.manifest, profiles.capability, profiles.mapping, profiles.namespace_registry)
    assert result.generation is None
    assert any(item["status"] == "blocked" for item in result.audit["findings"])


def test_approved_editorial_range_exception_completes_only_a_missing_range():
    profiles = load_generation_profiles(FIXTURES / "import-manifest-missing-range.yaml")
    mapping = dict(profiles.mapping)
    mapping["editorial_exceptions"] = [{"id": "range", "resource_uri": "https://example.org/source/property", "field": "hasRange", "reference_uri": "https://example.org/source/Parent", "status": "approved", "rationale": "Source omission.", "approved_by": "Editor", "approved_at": "2026-09-23", "decision_reference": "decision-1"}]
    inventory = load_inventory(FIXTURES / "rdf/missing-range.ttl", "turtle")
    result = resolve_generation(inventory, profiles.manifest, profiles.capability, mapping, profiles.namespace_registry)
    assert result.generation is not None
    xml, trace = write_xml(result.generation, profiles.capability, verify_capability_xsd(profiles.capability), inventory.source_sha256)
    assert etree.fromstring(xml).xpath("/namespace/properties/property[identifierInNamespace='property']/hasRange/text()") == ["Parent"]
    assert any(entry["origin"] == "editorial_exception" and entry["exception_id"] == "range" for entry in trace["entries"])


def test_generate_cli_writes_only_the_phase4_artifacts(tmp_path):
    assert main(["generate", "--manifest", str(FIXTURES / "import-manifest.yaml"), "--output-dir", str(tmp_path)]) == 0
    assert {path.name for path in tmp_path.iterdir()} == {"import.xml", "generation-trace.json", "generation-audit.json"}


def test_generate_cli_blocks_without_publishing_xml(tmp_path):
    assert main(["generate", "--manifest", str(FIXTURES / "import-manifest-missing-range.yaml"), "--output-dir", str(tmp_path)]) == 3
    assert {path.name for path in tmp_path.iterdir()} == {"generation-audit.json"}


def test_generate_blocker_annotates_an_optional_workbook(tmp_path):
    workbook = tmp_path / "mapping.xlsx"
    assert main(["assist", "export", "--manifest", str(FIXTURES / "import-manifest-missing-range.yaml"), "--output", str(workbook)]) == 0
    assert main([
        "generate", "--manifest", str(FIXTURES / "import-manifest-missing-range.yaml"),
        "--output-dir", str(tmp_path / "import"), "--workbook", str(workbook),
    ]) == 3
    validation = load_workbook(workbook)["VALIDATION"]
    assert any(row[0] == "generation" and row[1] == "error" for row in validation.iter_rows(min_row=2, values_only=True))


def test_generate_refuses_to_replace_an_existing_bundle(tmp_path):
    assert main(["generate", "--manifest", str(FIXTURES / "import-manifest.yaml"), "--output-dir", str(tmp_path)]) == 0
    original_xml = (tmp_path / "import.xml").read_bytes()
    assert main(["generate", "--manifest", str(FIXTURES / "import-manifest-missing-range.yaml"), "--output-dir", str(tmp_path)]) == 2
    assert (tmp_path / "import.xml").read_bytes() == original_xml


def test_generate_cli_is_deterministic_across_output_directories(tmp_path):
    first = tmp_path / "first"
    second = tmp_path / "second"
    command = ["generate", "--manifest", str(FIXTURES / "import-manifest.yaml"), "--output-dir"]
    assert main([*command, str(first)]) == 0
    assert main([*command, str(second)]) == 0
    for name in ("import.xml", "generation-trace.json", "generation-audit.json"):
        assert (first / name).read_bytes() == (second / name).read_bytes()


def test_deprecated_external_namespace_is_generated_and_audited():
    profiles = load_generation_profiles(FIXTURES / "import-manifest.yaml")
    registry = dict(profiles.namespace_registry)
    registry["namespaces"] = [{**registry["namespaces"][0], "status": "deprecated"}]
    inventory = load_inventory(FIXTURES / "rdf/valid.ttl", "turtle")
    result = resolve_generation(inventory, profiles.manifest, profiles.capability, profiles.mapping, registry)
    assert result.generation is not None
    assert any(item.get("reason") == "Deprecated external namespace used." for item in result.audit["findings"])


def test_generation_uses_a_generic_external_reference_rule_and_traces_it():
    profiles = load_generation_profiles(FIXTURES / "import-manifest.yaml")
    mapping = dict(profiles.mapping)
    mapping["external_references"] = []
    mapping["external_reference_rules"] = [{"id": "external-suffix", "uri_prefix": "https://example.org/external/", "reference_namespace": 100, "identifier_extraction": {"source": "uri_suffix"}}]
    inventory = load_inventory(FIXTURES / "rdf/valid.ttl", "turtle")
    result = resolve_generation(inventory, profiles.manifest, profiles.capability, mapping, profiles.namespace_registry)
    assert result.generation is not None
    xml, trace = write_xml(result.generation, profiles.capability, verify_capability_xsd(profiles.capability), inventory.source_sha256)
    assert etree.fromstring(xml).xpath("/namespace/classes/class[1]/equivalentClass/text()") == ["ExternalClass"]
    assert {entry["reference_rule"] for entry in trace["entries"] if entry["origin"] == "external_rule"} == {"external-suffix"}


def test_generation_uses_a_configured_literal_identifier_and_traces_its_triple(tmp_path):
    source = tmp_path / "literal-id.ttl"
    source.write_text(
        """@prefix ex: <https://example.org/source/> .
@prefix owl: <http://www.w3.org/2002/07/owl#> .
@prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
ex:DescriptiveClass a owl:Class ; ex:code "C1" ; rdfs:label "Class"@en .
""",
        encoding="utf-8",
    )
    profiles = load_generation_profiles(FIXTURES / "import-manifest.yaml")
    mapping = dict(profiles.mapping)
    mapping["rules"] = [{
        "id": "class", "selector": {"rdf_type": "http://www.w3.org/2002/07/owl#Class"}, "action": "map",
        "target": {"entity_kind": "class", "identifier_in_namespace": {"source": "literal_predicate", "predicate": "https://example.org/source/code"}, "label_predicates": ["http://www.w3.org/2000/01/rdf-schema#label"]},
    }]
    inventory = load_inventory(source, "turtle")
    result = resolve_generation(inventory, profiles.manifest, profiles.capability, mapping, profiles.namespace_registry)
    assert result.generation is not None
    xml, trace = write_xml(result.generation, profiles.capability, verify_capability_xsd(profiles.capability), inventory.source_sha256)
    assert etree.fromstring(xml).xpath("/namespace/classes/class/identifierInNamespace/text()") == ["C1"]
    identifier_entry = next(entry for entry in trace["entries"] if entry["element"] == "identifierInNamespace")
    assert identifier_entry["origin"] == "rdf" and len(identifier_entry["source_triples"]) == 1


def test_generation_uses_a_configured_regex_identifier():
    profiles = load_generation_profiles(FIXTURES / "import-manifest.yaml")
    mapping = dict(profiles.mapping)
    mapping["rules"] = [dict(rule) for rule in profiles.mapping["rules"]]
    mapping["rules"][0]["target"] = dict(mapping["rules"][0]["target"])
    mapping["rules"][0]["target"]["identifier_in_namespace"] = {"source": "regex_capture", "pattern": r"https://example\.org/source/([A-Za-z]+)"}
    result = resolve_generation(load_inventory(FIXTURES / "rdf/valid.ttl", "turtle"), profiles.manifest, profiles.capability, mapping, profiles.namespace_registry)
    assert result.generation is not None
    assert result.generation.classes[0].identifier == "Class"


def test_generation_orders_text_fields_per_xsd_and_writes_disjointness(tmp_path):
    source = tmp_path / "relations.ttl"
    source.write_text(
        """@prefix ex: <https://example.org/source/> .
@prefix owl: <http://www.w3.org/2002/07/owl#> .
@prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
@prefix skos: <http://www.w3.org/2004/02/skos/core#> .
ex:Alpha a owl:Class ; rdfs:label "Alpha"@en ; owl:disjointWith ex:Beta ; skos:scopeNote "Scope"@en ; skos:example "Example"@en ; rdfs:comment "Context"@en ; ex:bibliography "Bibliography"@en .
ex:Beta a owl:Class ; rdfs:label "Beta"@en .
""",
        encoding="utf-8",
    )
    profiles = load_generation_profiles(FIXTURES / "import-manifest.yaml")
    capability = {**profiles.capability, "constructs": {**profiles.capability["constructs"], "scope_note": "supported", "example": "supported"}}
    mapping = dict(profiles.mapping)
    mapping["rules"] = [{
        "id": "class", "selector": {"rdf_type": "http://www.w3.org/2002/07/owl#Class"}, "action": "map",
        "target": {"entity_kind": "class", "identifier_in_namespace": {"source": "uri_suffix", "strip_prefix": "https://example.org/source/"}, "label_predicates": ["http://www.w3.org/2000/01/rdf-schema#label"], "relations": [{"field": "disjointWith", "predicate": "http://www.w3.org/2002/07/owl#disjointWith"}], "text_fields": [{"field": "scopeNote", "predicates": ["http://www.w3.org/2004/02/skos/core#scopeNote"]}, {"field": "example", "predicates": ["http://www.w3.org/2004/02/skos/core#example"]}, {"field": "contextNote", "predicates": ["http://www.w3.org/2000/01/rdf-schema#comment"]}, {"field": "bibliographicalNote", "predicates": ["https://example.org/source/bibliography"]}]},
    }]
    inventory = load_inventory(source, "turtle")
    result = resolve_generation(inventory, profiles.manifest, capability, mapping, profiles.namespace_registry)
    assert result.generation is not None
    xml, _ = write_xml(result.generation, capability, verify_capability_xsd(capability), inventory.source_sha256)
    document = etree.fromstring(xml)
    assert document.xpath("/namespace/classes/class[identifierInNamespace='Alpha']/disjointWith/text()") == ["Beta"]
    assert [element.tag for element in document.xpath("/namespace/classes/class[identifierInNamespace='Alpha']/textProperties/*")] == ["scopeNote", "example", "contextNote", "bibliographicalNote"]
